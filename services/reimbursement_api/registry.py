"""Per-deal one-pager: capacity, packet history, open structuring proposals.

Returns a JSON snapshot AND writes a Word doc the principal can pull at
any time. Intended cadence: weekly during active reimbursement cycles,
on-demand otherwise.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from services.common.config import REPO_ROOT
from services.common.db import cursor


@dataclass
class RegistrySnapshot:
    deal_slug: str
    deal_name: str
    instruments: list[dict[str, Any]]
    open_proposals: list[dict[str, Any]]
    word_path: str


def _gather(deal_slug: str) -> dict[str, Any]:
    with cursor() as cur:
        cur.execute(
            "SELECT id, slug, name, jurisdiction, county, state, status FROM deals WHERE slug = %s",
            (deal_slug,),
        )
        deal = cur.fetchone()
        if deal is None:
            raise RuntimeError(f"Deal {deal_slug!r} not found.")
        deal = dict(deal)

        cur.execute(
            """
            SELECT i.id, i.kind::text AS kind, i.name, i.status,
                   i.authorized_amount, i.issued_amount, i.capacity_remaining,
                   COALESCE(
                       (SELECT SUM(eligible_amount)
                        FROM ledger_classifications lc
                        WHERE lc.instrument_id = i.id
                          AND lc.advisor_review IN ('pending', 'approved')),
                       0
                   ) AS classified_eligible,
                   COALESCE(
                       (SELECT SUM(requested_amount) FROM reimbursements r WHERE r.instrument_id = i.id),
                       0
                   ) AS total_requested,
                   COALESCE(
                       (SELECT SUM(approved_amount) FROM reimbursements r WHERE r.instrument_id = i.id),
                       0
                   ) AS total_approved,
                   COALESCE(
                       (SELECT SUM(paid_amount) FROM reimbursements r WHERE r.instrument_id = i.id),
                       0
                   ) AS total_paid,
                   (SELECT COUNT(*) FROM reimbursements r
                    WHERE r.instrument_id = i.id AND r.status IN ('drafting', 'submitted')) AS open_packets,
                   (SELECT COUNT(*) FROM reimbursements r
                    WHERE r.instrument_id = i.id AND r.status = 'paid') AS paid_packets
            FROM instruments i
            WHERE i.deal_id = %s
            ORDER BY i.kind::text
            """,
            (deal["id"],),
        )
        instruments = [dict(r) for r in cur.fetchall()]

        cur.execute(
            """
            SELECT id, family, title, estimated_dollars,
                   bond_counsel_view, municipal_advisor_view, pid_admin_view,
                   incorporated_into, created_at
            FROM pf_positions
            WHERE deal_id = %s
              AND (bond_counsel_view = 'pending'
                   OR municipal_advisor_view = 'pending'
                   OR pid_admin_view = 'pending')
            ORDER BY estimated_dollars DESC NULLS LAST, created_at DESC
            LIMIT 10
            """,
            (deal["id"],),
        )
        open_proposals = [dict(r) for r in cur.fetchall()]

    return {"deal": deal, "instruments": instruments, "open_proposals": open_proposals}


def _fmt_money(v: Any) -> str:
    if v is None:
        return "(n/a)"
    try:
        return f"${float(v):,.2f}"
    except (TypeError, ValueError):
        return str(v)


def _build_word(snap: dict[str, Any], out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    deal = snap["deal"]
    out_path = out_dir / f"{deal['slug']}_registry_{datetime.utcnow():%Y%m%d_%H%M}.docx"

    doc = Document()
    doc.add_heading(
        f"Reimbursement registry — {deal['name']}",
        level=0,
    )
    doc.add_paragraph(
        f"Jurisdiction: {deal['jurisdiction']}, {deal['county']} County, {deal['state']}\n"
        f"Status: {deal['status']}\n"
        f"Generated: {datetime.utcnow():%Y-%m-%d %H:%M UTC}"
    )

    doc.add_heading("Instruments", level=1)
    table = doc.add_table(rows=1, cols=8)
    table.style = "Light List"
    hdr = table.rows[0].cells
    for i, h in enumerate(
        [
            "Kind",
            "Name",
            "Status",
            "Authorized",
            "Issued",
            "Capacity remaining",
            "Classified (pend+app)",
            "Requested / approved / paid",
        ]
    ):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True

    for i in snap["instruments"]:
        row = table.add_row().cells
        row[0].text = i["kind"]
        row[1].text = i["name"]
        row[2].text = i["status"]
        row[3].text = _fmt_money(i["authorized_amount"])
        row[4].text = _fmt_money(i["issued_amount"])
        row[5].text = _fmt_money(i["capacity_remaining"])
        row[6].text = _fmt_money(i["classified_eligible"])
        row[7].text = (
            f"{_fmt_money(i['total_requested'])} / "
            f"{_fmt_money(i['total_approved'])} / "
            f"{_fmt_money(i['total_paid'])} "
            f"({i['open_packets']} open, {i['paid_packets']} paid)"
        )

    doc.add_heading("Top open structuring proposals", level=1)
    if not snap["open_proposals"]:
        doc.add_paragraph("(no open proposals — every position has all three advisor views recorded)")
    else:
        doc.add_paragraph(
            "Proposals where one or more advisors have not yet recorded a view. "
            "These are the highest-leverage open items for the next advisor session."
        )
        for p in snap["open_proposals"]:
            para = doc.add_paragraph()
            run = para.add_run(p["title"])
            run.bold = True
            if (p.get("estimated_dollars") or 0) >= 1_000_000:
                run.font.color.rgb = RGBColor(192, 0, 0)
            run.font.size = Pt(11)
            para.add_run(
                f"\n  Family: {p['family']}  Estimate: {_fmt_money(p['estimated_dollars'])}\n"
                f"  Bond counsel: {p['bond_counsel_view']}  "
                f"Municipal advisor: {p['municipal_advisor_view']}  "
                f"PID admin: {p['pid_admin_view']}"
            ).font.size = Pt(10)

    doc.save(out_path)
    return out_path


def render_registry(deal_slug: str) -> RegistrySnapshot:
    snap = _gather(deal_slug)
    out_dir = REPO_ROOT / "templates" / "deliverables" / "registry"
    word_path = _build_word(snap, out_dir)

    return RegistrySnapshot(
        deal_slug=snap["deal"]["slug"],
        deal_name=snap["deal"]["name"],
        instruments=[
            {
                "kind": i["kind"],
                "name": i["name"],
                "status": i["status"],
                "authorized_amount": float(i["authorized_amount"]) if i["authorized_amount"] is not None else None,
                "issued_amount": float(i["issued_amount"]) if i["issued_amount"] is not None else None,
                "capacity_remaining": float(i["capacity_remaining"]) if i["capacity_remaining"] is not None else None,
                "classified_eligible": float(i["classified_eligible"]),
                "total_requested": float(i["total_requested"]),
                "total_approved": float(i["total_approved"]),
                "total_paid": float(i["total_paid"]),
                "open_packets": i["open_packets"],
                "paid_packets": i["paid_packets"],
            }
            for i in snap["instruments"]
        ],
        open_proposals=[
            {
                "id": p["id"],
                "family": p["family"],
                "title": p["title"],
                "estimated_dollars": float(p["estimated_dollars"]) if p["estimated_dollars"] is not None else None,
                "bond_counsel_view": p["bond_counsel_view"],
                "municipal_advisor_view": p["municipal_advisor_view"],
                "pid_admin_view": p["pid_admin_view"],
                "incorporated_into": p["incorporated_into"],
            }
            for p in snap["open_proposals"]
        ],
        word_path=str(word_path),
    )
