"""Reimbursement packet assembly.

Pulls every classified-and-eligible line for the target instrument, groups
by category, generates a Word cover memo + Excel exhibits, builds the
missing-backup punch list, and writes a `reimbursements` row with
status='drafting' so the firm has a single source of truth for what
shipped to which advisor.

The Word cover carries a "COUNSEL MUST OPINE" banner — every packet is
a draft until bond counsel + municipal advisor + PID administrator
sign off. The Excel workbook is the per-line ground truth advisors
review.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services.common.config import REPO_ROOT
from services.common.db import cursor
from services.reimbursement_api.schemas import PFInstrumentKind


@dataclass
class PacketAssembly:
    deal_slug: str
    instrument_kind: str
    instrument_name: str
    packet_number: str
    reimbursement_id: str
    n_lines: int
    n_categories: int
    requested_total: float
    period_start: str | None
    period_end: str | None
    word_path: str
    excel_path: str
    punch_list: list[dict[str, Any]]


def _resolve_instrument(deal_slug: str, kind: str) -> dict[str, Any]:
    with cursor() as cur:
        cur.execute(
            """
            SELECT i.id, i.name, i.kind::text AS kind, i.authorized_amount,
                   i.issued_amount, i.capacity_remaining,
                   d.id AS deal_id, d.name AS deal_name, d.slug AS deal_slug
            FROM instruments i
            JOIN deals d ON d.id = i.deal_id
            WHERE d.slug = %s AND i.kind = %s::pf_instrument_kind
            """,
            (deal_slug, kind),
        )
        row = cur.fetchone()
    if row is None:
        raise RuntimeError(
            f"No instrument of kind {kind!r} found for deal {deal_slug!r}. "
            "Run `python -m services.reimbursement_api seed` first."
        )
    return dict(row)


def _gather_lines(instrument_id: str) -> list[dict[str, Any]]:
    sql = """
        SELECT lc.id AS classification_id, lc.category, lc.eligible_amount,
               lc.confidence, lc.rationale, lc.advisor_review,
               cl.id AS cost_ledger_id, cl.posted_on, cl.vendor,
               cl.invoice_number, cl.description, cl.amount,
               cl.gl_account, cl.backup_uri
        FROM ledger_classifications lc
        JOIN cost_ledger cl ON cl.id = lc.cost_ledger_id
        WHERE lc.instrument_id = %s
          AND lc.advisor_review IN ('pending', 'approved')
          AND lc.eligible_amount > 0
        ORDER BY lc.category ASC, cl.posted_on ASC
    """
    with cursor() as cur:
        cur.execute(sql, (instrument_id,))
        return [dict(r) for r in cur.fetchall()]


def _build_word(
    *,
    instrument: dict[str, Any],
    packet_number: str,
    lines: list[dict[str, Any]],
    by_category: dict[str, list[dict[str, Any]]],
    punch_list: list[dict[str, Any]],
    period_start: str | None,
    period_end: str | None,
    requested_total: float,
    out_dir: Path,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / (
        f"{instrument['deal_slug']}_{instrument['kind']}_packet_{packet_number}.docx"
    )

    doc = Document()
    doc.add_heading(
        f"Reimbursement Packet Draft — {instrument['deal_name']} — {instrument['name']}",
        level=0,
    )
    doc.add_paragraph(
        f"Packet number: {packet_number}\n"
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}\n"
        f"Period: {period_start or 'n/a'} to {period_end or 'n/a'}\n"
        f"Requested total: ${requested_total:,.2f}"
    )

    banner = doc.add_paragraph()
    run = banner.add_run(
        "COUNSEL MUST OPINE — Every classification below is a draft. Bond counsel "
        "decides eligibility; the municipal advisor decides what's market; the PID "
        "administrator decides operational defensibility. Do not submit until "
        "advisor sign-offs are recorded in pf_positions / ledger_classifications."
    )
    run.bold = True
    run.font.color.rgb = RGBColor(192, 0, 0)
    run.font.size = Pt(10)

    doc.add_heading("Summary by category", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light List"
    hdr = table.rows[0].cells
    hdr[0].text = "Category"
    hdr[1].text = "Lines"
    hdr[2].text = "Eligible amount"
    for cat in sorted(by_category):
        cat_lines = by_category[cat]
        cat_total = sum(float(l["eligible_amount"]) for l in cat_lines)
        row = table.add_row().cells
        row[0].text = cat
        row[1].text = str(len(cat_lines))
        row[2].text = f"${cat_total:,.2f}"
    total_row = table.add_row().cells
    total_row[0].text = "TOTAL"
    total_row[1].text = str(len(lines))
    total_row[2].text = f"${requested_total:,.2f}"

    doc.add_heading("Missing-backup punch list", level=1)
    if punch_list:
        doc.add_paragraph(
            f"{len(punch_list)} lines lack `backup_uri`. Attach the invoice / pay app / "
            "supporting evidence before the advisor review or the packet will be returned."
        )
        for p in punch_list:
            doc.add_paragraph(
                f"• {p['posted_on']} — {p['vendor'] or '(no vendor)'} — "
                f"{p['description']} — ${float(p['amount']):,.2f}\n"
                f"  Invoice: {p.get('invoice_number') or '(none)'}\n"
                f"  Category: {p['category']}"
            )
    else:
        doc.add_paragraph("(none — every classified line has a backup_uri)")

    doc.add_heading("Eligibility narrative by category", level=1)
    for cat in sorted(by_category):
        cat_lines = by_category[cat]
        doc.add_heading(cat, level=2)
        for line in cat_lines:
            p = doc.add_paragraph()
            p.add_run(
                f"{line['posted_on']} — {line['vendor'] or '(no vendor)'} — "
                f"${float(line['eligible_amount']):,.2f}"
            ).bold = True
            p.add_run(
                f"\n  Description: {line['description']}\n"
                f"  Rationale: {line['rationale']}\n"
                f"  Confidence: {line['confidence']:.2f}  "
                f"  Advisor review: {line['advisor_review']}\n"
                f"  Backup: {line.get('backup_uri') or '(MISSING)'}"
            )

    doc.save(out_path)
    return out_path


def _build_excel(
    *,
    instrument: dict[str, Any],
    packet_number: str,
    lines: list[dict[str, Any]],
    by_category: dict[str, list[dict[str, Any]]],
    requested_total: float,
    out_dir: Path,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / (
        f"{instrument['deal_slug']}_{instrument['kind']}_packet_{packet_number}.xlsx"
    )

    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    bold = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFFFE599")

    summary["A1"] = "Reimbursement Packet"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Deal: {instrument['deal_name']}"
    summary["A3"] = f"Instrument: {instrument['name']} ({instrument['kind']})"
    summary["A4"] = f"Packet number: {packet_number}"
    summary["A5"] = f"Generated: {datetime.utcnow():%Y-%m-%d %H:%M UTC}"

    summary["A7"] = "Category"
    summary["B7"] = "Lines"
    summary["C7"] = "Eligible amount"
    for c in ("A7", "B7", "C7"):
        summary[c].font = bold
        summary[c].fill = yellow

    r = 8
    for cat in sorted(by_category):
        cat_lines = by_category[cat]
        cat_total = sum(float(l["eligible_amount"]) for l in cat_lines)
        summary.cell(r, 1, cat)
        summary.cell(r, 2, len(cat_lines))
        summary.cell(r, 3, cat_total).number_format = "$#,##0.00"
        r += 1
    summary.cell(r, 1, "TOTAL").font = bold
    summary.cell(r, 2, len(lines)).font = bold
    cell = summary.cell(r, 3, requested_total)
    cell.font = bold
    cell.number_format = "$#,##0.00"

    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 8
    summary.column_dimensions["C"].width = 18

    detail = wb.create_sheet("Line Items")
    headers = [
        "Category",
        "Posted on",
        "Vendor",
        "Invoice #",
        "Description",
        "Line amount",
        "Eligible amount",
        "Confidence",
        "Advisor review",
        "Rationale",
        "Backup URI",
        "GL account",
    ]
    detail.append(headers)
    for cell_idx in range(1, len(headers) + 1):
        c = detail.cell(1, cell_idx)
        c.font = bold
        c.fill = yellow
        c.alignment = Alignment(wrap_text=True)

    for line in lines:
        detail.append(
            [
                line["category"],
                str(line["posted_on"]),
                line.get("vendor") or "",
                line.get("invoice_number") or "",
                line["description"],
                float(line["amount"]),
                float(line["eligible_amount"]),
                round(float(line["confidence"] or 0), 3),
                line["advisor_review"],
                line["rationale"],
                line.get("backup_uri") or "",
                line.get("gl_account") or "",
            ]
        )

    for col_letter, width in zip(
        ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"],
        [22, 12, 16, 14, 60, 14, 14, 10, 14, 60, 30, 14],
    ):
        detail.column_dimensions[col_letter].width = width

    for row in detail.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = "$#,##0.00"

    wb.save(out_path)
    return out_path


def _insert_reimbursement(
    *,
    instrument_id: str,
    packet_number: str,
    requested_total: float,
    packet_uri: str,
) -> str:
    with cursor() as cur:
        cur.execute(
            """
            INSERT INTO reimbursements (
                instrument_id, packet_number, requested_amount, status, packet_uri
            )
            VALUES (%s, %s, %s, 'drafting', %s)
            ON CONFLICT (instrument_id, packet_number) DO UPDATE SET
                requested_amount = EXCLUDED.requested_amount,
                packet_uri = EXCLUDED.packet_uri,
                status = 'drafting'
            RETURNING id
            """,
            (instrument_id, packet_number, requested_total, packet_uri),
        )
        return cur.fetchone()["id"]


def assemble_packet(
    *,
    deal_slug: str,
    instrument_kind: PFInstrumentKind,
    packet_number: str | None = None,
) -> PacketAssembly:
    if packet_number is None:
        packet_number = datetime.utcnow().strftime("%Y%m%d-%H%M")

    instrument = _resolve_instrument(deal_slug, instrument_kind)
    lines = _gather_lines(instrument["id"])
    if not lines:
        raise RuntimeError(
            f"No pending/approved classifications for {instrument['name']!r}. "
            "Run `python -m services.reimbursement_api classify` first."
        )

    by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for line in lines:
        by_category[line["category"]].append(line)

    punch_list = [l for l in lines if not l.get("backup_uri")]
    requested_total = float(sum(Decimal(str(l["eligible_amount"])) for l in lines))
    posted = sorted(str(l["posted_on"]) for l in lines if l.get("posted_on"))
    period_start = posted[0] if posted else None
    period_end = posted[-1] if posted else None

    out_dir = REPO_ROOT / "templates" / "deliverables" / "packets"
    word_path = _build_word(
        instrument=instrument,
        packet_number=packet_number,
        lines=lines,
        by_category=by_category,
        punch_list=punch_list,
        period_start=period_start,
        period_end=period_end,
        requested_total=requested_total,
        out_dir=out_dir,
    )
    excel_path = _build_excel(
        instrument=instrument,
        packet_number=packet_number,
        lines=lines,
        by_category=by_category,
        requested_total=requested_total,
        out_dir=out_dir,
    )

    reimbursement_id = _insert_reimbursement(
        instrument_id=instrument["id"],
        packet_number=packet_number,
        requested_total=requested_total,
        packet_uri=str(word_path),
    )

    return PacketAssembly(
        deal_slug=deal_slug,
        instrument_kind=instrument_kind,
        instrument_name=instrument["name"],
        packet_number=packet_number,
        reimbursement_id=reimbursement_id,
        n_lines=len(lines),
        n_categories=len(by_category),
        requested_total=requested_total,
        period_start=period_start,
        period_end=period_end,
        word_path=str(word_path),
        excel_path=str(excel_path),
        punch_list=[
            {
                "posted_on": str(p["posted_on"]),
                "vendor": p.get("vendor"),
                "invoice_number": p.get("invoice_number"),
                "description": p["description"],
                "amount": float(p["amount"]),
                "category": p["category"],
            }
            for p in punch_list
        ],
    )
