"""Daily and weekly Lavon-watch digests.

Daily: list every signal from the last 24h ranked by impact_score,
rendered as a Word doc the operator can drop into the deal folder.

Weekly: Opus 4.7 synthesizes the past 7 days against the pilot context —
"what changed about our political/entitlement posture this week?" — and
the synthesis goes into the long-form weekly Word doc.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from services.common.anthropic_client import OPUS, client
from services.common.config import REPO_ROOT
from services.common.db import cursor
from services.signal_collectors.pilot_context import load_pilot_context_text


WEEKLY_SYNTHESIS_SYSTEM = """You are the Lavon-watch weekly synthesist. Read the past 7 days of signals against the active pilot context and produce a single coherent synthesis.

Output sections, in order:
1. POSTURE SHIFT — 2-3 sentences. What changed this week about our political and entitlement posture? Include direction (better/worse/neutral) and confidence.
2. CRITICAL-PATH IMPACTS — bullet list. Each bullet ties one or more signals (cite by title) to a specific instrument or critical-path event.
3. EMERGING WATCHES — bullet list. Signals that aren't yet critical but warrant tighter tracking.
4. WHAT TO ASK — bullet list of specific questions the principal should put to advisors, council, or staff this week, anchored to the signals.

Do not invent context. Cite signals by their title in [brackets]. If no signals were ingested this week, say so plainly — do not pad.
"""


@dataclass
class DigestRow:
    title: str
    summary: str
    kind: str
    source: str
    source_url: str
    occurred_on: date | None
    impact_score: int
    impact_rationale: str
    affected_instruments: list[str]


def _gather_period_signals(
    *, period_start: datetime, period_end: datetime
) -> list[DigestRow]:
    sql = """
        SELECT s.id, s.kind::text AS kind, s.source::text AS source,
               s.title, s.summary, s.source_url, s.occurred_on,
               s.impact_score, s.impact_rationale,
               COALESCE(
                   (
                       SELECT array_agg(DISTINCT i.kind::text ORDER BY i.kind::text)
                       FROM signal_pilot_links spl
                       JOIN instruments i ON i.id = spl.instrument_id
                       WHERE spl.signal_id = s.id
                   ),
                   ARRAY[]::text[]
               ) AS affected_instruments
        FROM signals s
        WHERE s.retrieved_at >= %s AND s.retrieved_at < %s
        ORDER BY s.impact_score DESC, s.occurred_on DESC NULLS LAST, s.retrieved_at DESC
    """
    with cursor() as cur:
        cur.execute(sql, (period_start, period_end))
        rows = cur.fetchall()
    return [
        DigestRow(
            title=r["title"],
            summary=r["summary"],
            kind=r["kind"],
            source=r["source"],
            source_url=r["source_url"],
            occurred_on=r["occurred_on"],
            impact_score=r["impact_score"],
            impact_rationale=r["impact_rationale"],
            affected_instruments=list(r["affected_instruments"] or []),
        )
        for r in rows
    ]


def _add_signal(doc: Document, row: DigestRow) -> None:
    h = doc.add_paragraph()
    title_run = h.add_run(f"[{row.impact_score}] {row.title}")
    title_run.bold = True
    title_run.font.size = Pt(11)
    if row.impact_score >= 4:
        title_run.font.color.rgb = RGBColor(192, 0, 0)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"kind={row.kind}  source={row.source}  "
        f"date={row.occurred_on or 'n/a'}  "
        f"instruments={','.join(row.affected_instruments) or '(none)'}\n"
        f"{row.source_url}"
    )
    meta_run.font.size = Pt(9)

    summary_p = doc.add_paragraph(row.summary)
    summary_p.runs[0].font.size = Pt(10)

    rat_p = doc.add_paragraph()
    rat_p.add_run("Impact rationale: ").bold = True
    rat_p.add_run(row.impact_rationale).font.size = Pt(10)


def render_daily(*, period_start: datetime, period_end: datetime) -> dict[str, Any]:
    rows = _gather_period_signals(period_start=period_start, period_end=period_end)

    out_dir = REPO_ROOT / "templates" / "deliverables" / "lavon_daily"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"lavon_daily_{period_start:%Y%m%d}.docx"

    doc = Document()
    doc.add_heading(
        f"Lavon-watch — daily — {period_start:%Y-%m-%d}", level=0
    )
    doc.add_paragraph(
        f"Window: {period_start:%Y-%m-%d %H:%MZ} to {period_end:%Y-%m-%d %H:%MZ}\n"
        f"Signals: {len(rows)}"
    )

    if not rows:
        doc.add_paragraph("(no signals in window)")
    else:
        high = [r for r in rows if r.impact_score >= 4]
        mid = [r for r in rows if 2 <= r.impact_score <= 3]
        low = [r for r in rows if r.impact_score <= 1]

        if high:
            doc.add_heading("High impact (4–5)", level=1)
            for r in high:
                _add_signal(doc, r)
        if mid:
            doc.add_heading("Mid impact (2–3)", level=1)
            for r in mid:
                _add_signal(doc, r)
        if low:
            doc.add_heading("Informational (0–1)", level=1)
            for r in low:
                _add_signal(doc, r)

    doc.save(out_path)
    high_impact_count = sum(1 for r in rows if r.impact_score >= 4)

    with cursor() as cur:
        cur.execute(
            """
            INSERT INTO lavon_digests (
                cadence, period_start, period_end, digest_uri, signal_count, high_impact_count, sent_at
            )
            VALUES ('daily', %s, %s, %s, %s, %s, NOW())
            ON CONFLICT (cadence, period_start, period_end) DO UPDATE SET
                digest_uri = EXCLUDED.digest_uri,
                signal_count = EXCLUDED.signal_count,
                high_impact_count = EXCLUDED.high_impact_count,
                sent_at = NOW()
            RETURNING id
            """,
            (
                period_start.date(),
                period_end.date(),
                str(out_path),
                len(rows),
                high_impact_count,
            ),
        )
        digest_id = cur.fetchone()["id"]

    return {
        "digest_id": digest_id,
        "cadence": "daily",
        "signal_count": len(rows),
        "high_impact_count": high_impact_count,
        "digest_uri": str(out_path),
    }


def render_weekly(*, week_starting: date) -> dict[str, Any]:
    period_start = datetime.combine(week_starting, datetime.min.time())
    period_end = period_start + timedelta(days=7)
    rows = _gather_period_signals(period_start=period_start, period_end=period_end)
    ctx = load_pilot_context_text()

    if rows:
        signal_block = "\n\n".join(
            f"[{r.impact_score}] {r.title}\n"
            f"  source={r.source} occurred_on={r.occurred_on}\n"
            f"  summary: {r.summary}\n"
            f"  rationale: {r.impact_rationale}\n"
            f"  instruments: {','.join(r.affected_instruments) or '(none)'}"
            for r in rows
        )
    else:
        signal_block = "(no signals ingested this week)"

    user_msg = (
        f"Week of {week_starting:%Y-%m-%d}.\n"
        f"Signals from this week:\n\n{signal_block}\n\n"
        "Produce the weekly synthesis per the system instructions."
    )

    a = client()
    response = a.messages.create(
        model=OPUS,
        max_tokens=4000,
        thinking={"type": "adaptive"},
        system=[
            {"type": "text", "text": WEEKLY_SYNTHESIS_SYSTEM},
            {"type": "text", "text": ctx["context_text"], "cache_control": {"type": "ephemeral"}},
        ],
        messages=[{"role": "user", "content": user_msg}],
    )
    synthesis = next((b.text for b in response.content if b.type == "text"), "")

    out_dir = REPO_ROOT / "templates" / "deliverables" / "lavon_weekly"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"lavon_weekly_{week_starting:%Y%m%d}.docx"

    doc = Document()
    doc.add_heading(f"Lavon-watch — weekly — week of {week_starting:%Y-%m-%d}", level=0)
    doc.add_paragraph(
        f"Window: {period_start:%Y-%m-%d} to {period_end:%Y-%m-%d}\n"
        f"Signals: {len(rows)}"
    )

    doc.add_heading("Synthesis", level=1)
    for para in synthesis.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_heading("Underlying signals", level=1)
    if not rows:
        doc.add_paragraph("(none)")
    else:
        for r in rows:
            _add_signal(doc, r)

    doc.save(out_path)
    high_impact_count = sum(1 for r in rows if r.impact_score >= 4)

    with cursor() as cur:
        cur.execute(
            """
            INSERT INTO lavon_digests (
                cadence, period_start, period_end, digest_uri, signal_count, high_impact_count, sent_at
            )
            VALUES ('weekly', %s, %s, %s, %s, %s, NOW())
            ON CONFLICT (cadence, period_start, period_end) DO UPDATE SET
                digest_uri = EXCLUDED.digest_uri,
                signal_count = EXCLUDED.signal_count,
                high_impact_count = EXCLUDED.high_impact_count,
                sent_at = NOW()
            RETURNING id
            """,
            (
                period_start.date(),
                period_end.date(),
                str(out_path),
                len(rows),
                high_impact_count,
            ),
        )
        digest_id = cur.fetchone()["id"]

    return {
        "digest_id": digest_id,
        "cadence": "weekly",
        "signal_count": len(rows),
        "high_impact_count": high_impact_count,
        "digest_uri": str(out_path),
        "synthesis": synthesis,
    }
